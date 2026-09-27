import pandas as pd
from sqlalchemy import create_engine
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

from src.database import get_db_engine

def export_management_report():
    print("📊 Yönetim Raporu Veritabanından Çekiliyor...")
    engine = get_db_engine()
    
    # 1. VIEW Yapısından Verileri Çekme
    query = """
    SELECT 
        stat_date AS "Tarih",
        hotel_type AS "Otel Tipi",
        total_bookings AS "Toplam Rezervasyon",
        active_bookings AS "Başarılı Konaklama",
        canceled_bookings AS "İptal Sayısı",
        avg_daily_rate AS "Ortalama Oda Fiyatı (ADR)",
        total_revenue AS "Toplam Gelir"
    FROM view_daily_hotel_stats
    ORDER BY stat_date DESC;
    """
    
    df = pd.read_sql(query, engine)
    
    if df.empty:
        print("⚠️ Uyarı: Raporlanacak veri bulunamadı!")
        return

    # 2. Excel Rapor Dosyası Yolu
    output_dir = "reports"
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "Hotel_Revenue_Management_Report.xlsx")

    # 3. Openpyxl İle Şık Bir Excel Tasarımı Oluşturma
    print("🎨 Excel Dosyası Biçimlendiriliyor...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Günlük Yönetim Özeti"

    # Başlık Satırı Ekleme (Şirket/Rapor İsmi)
    ws.merge_cells('A1:G1')
    top_title = ws['A1']
    top_title.value = "OTEL GELİR YÖNETİMİ VE PERFORMANS RAPORU"
    top_title.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    top_title.fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid') # Koyu Mavi
    top_title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40

    # Tablo Kolon Başlıkları
    headers = list(df.columns)
    ws.append([]) # 2. Satır Boş
    ws.append(headers) # 3. Satıra başlıklar

    header_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 25

    # Verileri Ekleme ve Formatlama
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for row in df.itertuples(index=False):
        ws.append(list(row))
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = thin_border
            cell.font = Font(name='Calibri', size=10)

            # Tarih Formatı
            if col_num == 1:
                cell.alignment = Alignment(horizontal='center')
            # Otel Tipi
            elif col_num == 2:
                cell.alignment = Alignment(horizontal='left')
            # Sayısal Adetler (Binlik Ayraçlı)
            elif col_num in [3, 4, 5]:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')
            # Ortalama Fiyat (Euro/TL Para Birimi Formatı)
            elif col_num == 6:
                cell.number_format = '€#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            # Toplam Gelir
            elif col_num == 7:
                cell.number_format = '€#,##0.00'
                cell.alignment = Alignment(horizontal='right')

    # Kolon Genişliklerini İçeriğe Göre Otomatik Ayarlama
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # 1. satır birleştirildiği için genişlik hesabına katmayalım
            if cell.row == 1:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Dosyayı Kaydetme
    wb.save(file_path)
    print(f"🚀 Rapor Başarıyla Oluşturuldu: {file_path}")

if __name__ == "__main__":
    export_management_report()
